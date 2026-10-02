import re
import os
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """Set cell margins in dxa (1 pt = 20 dxa)."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_cell_shading(cell, color_hex):
    """Set cell background color."""
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
    tcPr.append(shd)

def set_cell_border(cell, **kwargs):
    """
    Set cell borders:
    kwargs: top, bottom, left, right
    values: dict(sz=12, val='single', color='FF0000', space='0')
    """
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = OxmlElement('w:tcBorders')
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        edge_data = kwargs.get(edge)
        if edge_data:
            tag = f'w:{edge}'
            element = OxmlElement(tag)
            element.set(qn('w:val'), edge_data.get('val', 'single'))
            element.set(qn('w:sz'), str(edge_data.get('sz', 4)))
            element.set(qn('w:space'), str(edge_data.get('space', 0)))
            element.set(qn('w:color'), edge_data.get('color', 'auto'))
            tcBorders.append(element)
    tcPr.append(tcBorders)

def add_styled_paragraph(doc, text="", style=None, space_after=6, space_before=0, line_spacing=1.15):
    p = doc.add_paragraph(style=style)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.line_spacing = line_spacing
    return p

def parse_inline_formatting(paragraph, text, base_font_name="Arial", base_font_size=Pt(11), default_color=None):
    """Parse bold, italic, code, links into docx runs."""
    # Pattern to match:
    # 1. `code`
    # 2. **bold**
    # 3. *italic*
    # 4. [text](url)
    token_pattern = re.compile(
        r'(`(?P<code>[^`]+)`)'
        r'|(\*\*\*(?P<bi>[^*]+)\*\*\*)'
        r'|(\*\*(?P<bold>[^*]+)\*\*)'
        r'|(\*(?P<italic>[^*]+)\*)'
        r'|(\[(?P<link_text>[^\]]+)\]\((?P<link_url>[^)]+)\))'
    )
    
    last_end = 0
    for match in token_pattern.finditer(text):
        start, end = match.span()
        if start > last_end:
            plain_text = text[last_end:start]
            run = paragraph.add_run(plain_text)
            run.font.name = base_font_name
            run.font.size = base_font_size
            if default_color:
                run.font.color.rgb = default_color

        if match.group('code'):
            code_text = match.group('code')
            run = paragraph.add_run(code_text)
            run.font.name = "Consolas"
            run.font.size = Pt(base_font_size.pt * 0.9)
            run.font.color.rgb = RGBColor(199, 37, 78)  # soft code red/magenta
            # highlight or shading
        elif match.group('bi'):
            bi_text = match.group('bi')
            run = paragraph.add_run(bi_text)
            run.font.name = base_font_name
            run.font.size = base_font_size
            run.bold = True
            run.italic = True
            if default_color:
                run.font.color.rgb = default_color
        elif match.group('bold'):
            bold_text = match.group('bold')
            run = paragraph.add_run(bold_text)
            run.font.name = base_font_name
            run.font.size = base_font_size
            run.bold = True
            if default_color:
                run.font.color.rgb = default_color
        elif match.group('italic'):
            italic_text = match.group('italic')
            run = paragraph.add_run(italic_text)
            run.font.name = base_font_name
            run.font.size = base_font_size
            run.italic = True
            if default_color:
                run.font.color.rgb = default_color
        elif match.group('link_text'):
            link_text = match.group('link_text')
            run = paragraph.add_run(link_text)
            run.font.name = base_font_name
            run.font.size = base_font_size
            run.font.color.rgb = RGBColor(26, 86, 219)  # link blue
            run.underline = True

        last_end = end

    if last_end < len(text):
        remaining = text[last_end:]
        run = paragraph.add_run(remaining)
        run.font.name = base_font_name
        run.font.size = base_font_size
        if default_color:
            run.font.color.rgb = default_color

def create_table_from_md(doc, table_lines):
    rows_data = []
    for line in table_lines:
        line_clean = line.strip()
        if not line_clean.startswith('|'):
            continue
        cells = [c.strip() for c in line_clean.split('|')[1:-1]]
        # Check if separator row
        if all(re.match(r'^:?-+:?$', c) for c in cells):
            continue
        rows_data.append(cells)
        
    if not rows_data:
        return
        
    num_rows = len(rows_data)
    num_cols = max(len(r) for r in rows_data)
    
    # Normalize row lengths
    for r in rows_data:
        while len(r) < num_cols:
            r.append('')
            
    table = doc.add_table(rows=num_rows, cols=num_cols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    
    # Format table
    border_color = "D3D3D3"
    border_spec = dict(val="single", sz=4, color=border_color)
    
    for r_idx, row in enumerate(table.rows):
        is_header = (r_idx == 0)
        for c_idx, cell in enumerate(row.cells):
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_margins(cell, top=140, bottom=140, left=180, right=180)
            set_cell_border(cell, top=border_spec, bottom=border_spec, left=border_spec, right=border_spec)
            
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.1
            
            cell_text = rows_data[r_idx][c_idx]
            
            if is_header:
                set_cell_shading(cell, "1F4E79")  # Navy Blue
                parse_inline_formatting(p, f"**{cell_text}**", base_font_size=Pt(10), default_color=RGBColor(255, 255, 255))
            else:
                if r_idx % 2 == 1:
                    set_cell_shading(cell, "F9FAFB")  # Subtle zebra
                else:
                    set_cell_shading(cell, "FFFFFF")
                parse_inline_formatting(p, cell_text, base_font_size=Pt(10))
                
    doc.add_paragraph().paragraph_format.space_after = Pt(6)

def convert_md_to_docx(md_path, docx_path):
    with open(md_path, 'r', encoding='utf-8') as f:
        content = f.read()

    doc = Document()
    
    # Page Setup: A4, 1 inch margins
    sections = doc.sections
    for section in sections:
        section.page_width = Inches(8.27)
        section.page_height = Inches(11.69)
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    # Base Colors
    NAVY = RGBColor(31, 78, 121)       # #1F4E79
    STEEL_BLUE = RGBColor(46, 91, 136) # #2E5B88
    DARK_BLUE = RGBColor(51, 102, 153) # #336699
    SLATE = RGBColor(70, 80, 95)
    DARK_TEXT = RGBColor(33, 37, 41)
    
    lines = content.split('\n')
    i = 0
    in_table = False
    table_lines = []
    
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()
        
        # Check table
        if stripped.startswith('|') and stripped.endswith('|'):
            table_lines.append(stripped)
            i += 1
            continue
        elif table_lines:
            create_table_from_md(doc, table_lines)
            table_lines = []

        if not stripped:
            i += 1
            continue

        # Horizontal Rule
        if stripped in ('---', '***', '___'):
            # subtle divider line
            p = add_styled_paragraph(doc, space_after=12, space_before=12)
            p_border = parse_xml(f'<w:pBdr {nsdecls("w")}><w:bottom w:val="single" w:sz="6" w:space="1" w:color="D3D3D3"/></w:pBdr>')
            p._p.get_or_add_pPr().append(p_border)
            i += 1
            continue

        # Blockquote (> ...)
        if stripped.startswith('>'):
            quote_text = stripped[1:].strip()
            # Multi-line quote gathering
            quotes = [quote_text]
            while i + 1 < len(lines) and lines[i+1].strip().startswith('>'):
                i += 1
                quotes.append(lines[i].strip()[1:].strip())
            
            # Create a 1x1 table for callout box
            tbl = doc.add_table(rows=1, cols=1)
            tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
            cell = tbl.rows[0].cells[0]
            cell.width = Inches(6.27)
            set_cell_margins(cell, top=140, bottom=140, left=200, right=200)
            set_cell_shading(cell, "F3F4F6")
            set_cell_border(cell, 
                            left=dict(val="single", sz=24, color="1F4E79"),
                            top=dict(val="none"),
                            bottom=dict(val="none"),
                            right=dict(val="none"))
            
            p_box = cell.paragraphs[0]
            p_box.paragraph_format.space_before = Pt(0)
            p_box.paragraph_format.space_after = Pt(0)
            p_box.paragraph_format.line_spacing = 1.15
            full_quote = "  \n".join(quotes)
            for q_idx, q_line in enumerate(quotes):
                if q_idx > 0:
                    p_box = cell.add_paragraph()
                    p_box.paragraph_format.space_before = Pt(2)
                    p_box.paragraph_format.space_after = Pt(2)
                    p_box.paragraph_format.line_spacing = 1.15
                parse_inline_formatting(p_box, q_line, base_font_size=Pt(10), default_color=SLATE)
                
            doc.add_paragraph().paragraph_format.space_after = Pt(6)
            i += 1
            continue

        # Headings
        if stripped.startswith('# '):
            p = add_styled_paragraph(doc, space_before=14, space_after=8)
            parse_inline_formatting(p, f"**{stripped[2:].strip()}**", base_font_size=Pt(18), default_color=NAVY)
            i += 1
            continue
        elif stripped.startswith('## '):
            p = add_styled_paragraph(doc, space_before=16, space_after=6)
            parse_inline_formatting(p, f"**{stripped[3:].strip()}**", base_font_size=Pt(14), default_color=STEEL_BLUE)
            i += 1
            continue
        elif stripped.startswith('### '):
            p = add_styled_paragraph(doc, space_before=12, space_after=4)
            parse_inline_formatting(p, f"**{stripped[4:].strip()}**", base_font_size=Pt(12), default_color=DARK_BLUE)
            i += 1
            continue
        elif stripped.startswith('#### '):
            p = add_styled_paragraph(doc, space_before=8, space_after=3)
            parse_inline_formatting(p, f"**{stripped[5:].strip()}**", base_font_size=Pt(11), default_color=DARK_TEXT)
            i += 1
            continue

        # Unordered Lists (- , * , + )
        ul_match = re.match(r'^(\s*)([-*+])\s+(.*)$', line)
        if ul_match:
            indent_spaces = len(ul_match.group(1))
            item_text = ul_match.group(3)
            p = add_styled_paragraph(doc, space_before=2, space_after=3)
            p.paragraph_format.left_indent = Inches(0.25 * (indent_spaces // 2 + 1))
            bullet_char = "• " if indent_spaces < 3 else "– "
            run_b = p.add_run(bullet_char)
            run_b.font.name = "Arial"
            run_b.font.size = Pt(11)
            run_b.font.color.rgb = DARK_BLUE
            parse_inline_formatting(p, item_text, base_font_size=Pt(11))
            i += 1
            continue

        # Ordered Lists (1. , 2. )
        ol_match = re.match(r'^(\s*)(\d+)\.\s+(.*)$', line)
        if ol_match:
            indent_spaces = len(ol_match.group(1))
            num = ol_match.group(2)
            item_text = ol_match.group(3)
            p = add_styled_paragraph(doc, space_before=3, space_after=3)
            p.paragraph_format.left_indent = Inches(0.25 * (indent_spaces // 2 + 1))
            run_num = p.add_run(f"{num}. ")
            run_num.font.name = "Arial"
            run_num.font.size = Pt(11)
            run_num.bold = True
            run_num.font.color.rgb = DARK_BLUE
            parse_inline_formatting(p, item_text, base_font_size=Pt(11))
            i += 1
            continue

        # Normal Paragraph
        p = add_styled_paragraph(doc, space_before=2, space_after=6)
        parse_inline_formatting(p, stripped, base_font_size=Pt(11))
        i += 1

    # End of document, check any remaining table
    if table_lines:
        create_table_from_md(doc, table_lines)

    # Save
    doc.save(docx_path)
    print(f"File saved successfully to: {docx_path}")

if __name__ == "__main__":
    src_md = r"c:\Users\HP\Downloads\Phân cụm động lượng TTCK\docs\M1_M2_PRESENTATION_GUIDE.md"
    dst_docx = r"C:\Users\HP\Downloads\M1_M2_PRESENTATION_GUIDE.docx"
    convert_md_to_docx(src_md, dst_docx)
